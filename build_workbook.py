import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

df = pd.read_csv("/home/claude/projects/retail-sales/raw_sales_data.csv", parse_dates=["date"])

wb = Workbook()

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="E8622A")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=14)
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---------------- Raw Data sheet ----------------
ws = wb.active
ws.title = "Raw Data"
headers = ["order_id","date","region","category","channel","units","unit_price","revenue","cost","profit","month"]
ws.append(headers)
for c in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=c)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center")

n = len(df)
for i, row in df.iterrows():
    r = i + 2
    ws.cell(row=r, column=1, value=int(row.order_id))
    ws.cell(row=r, column=2, value=row.date.to_pydatetime())
    ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
    ws.cell(row=r, column=3, value=row.region)
    ws.cell(row=r, column=4, value=row.category)
    ws.cell(row=r, column=5, value=row.channel)
    ws.cell(row=r, column=6, value=int(row.units))
    ws.cell(row=r, column=7, value=float(row.unit_price))
    ws.cell(row=r, column=7).number_format = "$#,##0.00"
    ws.cell(row=r, column=8, value=float(row.revenue))
    ws.cell(row=r, column=8).number_format = "$#,##0.00"
    ws.cell(row=r, column=9, value=float(row.cost))
    ws.cell(row=r, column=9).number_format = "$#,##0.00"
    ws.cell(row=r, column=10, value=f"=H{r}-I{r}")
    ws.cell(row=r, column=10).number_format = "$#,##0.00"
    ws.cell(row=r, column=11, value=f'=TEXT(B{r},"mmm-yy")')

for col, width in zip("ABCDEFGHIJK", [10,12,10,16,12,7,11,12,12,12,9]):
    ws.column_dimensions[col].width = width
ws.freeze_panes = "A2"
last_row = n + 1

# ---------------- Summary sheet ----------------
ws2 = wb.create_sheet("Summary")
ws2["A1"] = "Retail Sales Performance — Summary"
ws2["A1"].font = TITLE_FONT
ws2.merge_cells("A1:D1")

# Region table
ws2["A3"] = "Revenue & Margin by Region"
ws2["A3"].font = Font(name=FONT, bold=True, size=12, color="E8622A")
region_headers = ["Region", "Revenue", "Cost", "Profit", "Margin %"]
for c, h in enumerate(region_headers, start=1):
    cell = ws2.cell(row=4, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL

regions = ["North", "South", "East", "West"]
for i, region in enumerate(regions):
    r = 5 + i
    ws2.cell(row=r, column=1, value=region).font = BODY_FONT
    ws2.cell(row=r, column=2, value=f"=SUMIF('Raw Data'!C:C,A{r},'Raw Data'!H:H)").number_format = "$#,##0"
    ws2.cell(row=r, column=3, value=f"=SUMIF('Raw Data'!C:C,A{r},'Raw Data'!I:I)").number_format = "$#,##0"
    ws2.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = "$#,##0"
    ws2.cell(row=r, column=5, value=f"=IFERROR(D{r}/B{r},0)").number_format = "0.0%"
    for c in range(1, 6):
        ws2.cell(row=r, column=c).border = BORDER

# Category table
ws2["A11"] = "Revenue by Category"
ws2["A11"].font = Font(name=FONT, bold=True, size=12, color="E8622A")
cat_headers = ["Category", "Revenue", "Units Sold"]
for c, h in enumerate(cat_headers, start=1):
    cell = ws2.cell(row=12, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL

categories = ["Electronics","Apparel","Home & Kitchen","Beauty","Sports"]
for i, cat in enumerate(categories):
    r = 13 + i
    ws2.cell(row=r, column=1, value=cat).font = BODY_FONT
    ws2.cell(row=r, column=2, value=f"=SUMIF('Raw Data'!D:D,A{r},'Raw Data'!H:H)").number_format = "$#,##0"
    ws2.cell(row=r, column=3, value=f"=SUMIF('Raw Data'!D:D,A{r},'Raw Data'!F:F)").number_format = "#,##0"
    for c in range(1, 4):
        ws2.cell(row=r, column=c).border = BORDER

# Channel table
ws2["A19"] = "Revenue by Channel"
ws2["A19"].font = Font(name=FONT, bold=True, size=12, color="E8622A")
ch_headers = ["Channel", "Revenue", "Orders"]
for c, h in enumerate(ch_headers, start=1):
    cell = ws2.cell(row=20, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
channels = ["Online","In-Store","Mobile App"]
for i, ch in enumerate(channels):
    r = 21 + i
    ws2.cell(row=r, column=1, value=ch).font = BODY_FONT
    ws2.cell(row=r, column=2, value=f"=SUMIF('Raw Data'!E:E,A{r},'Raw Data'!H:H)").number_format = "$#,##0"
    ws2.cell(row=r, column=3, value=f"=COUNTIF('Raw Data'!E:E,A{r})").number_format = "#,##0"
    for c in range(1, 4):
        ws2.cell(row=r, column=c).border = BORDER

for col, width in zip("ABCDE", [16,14,14,14,10]):
    ws2.column_dimensions[col].width = width

# Region bar chart
chart1 = BarChart()
chart1.title = "Revenue by region"
chart1.style = 10
chart1.y_axis.title = "Revenue ($)"
data = Reference(ws2, min_col=2, min_row=4, max_row=8)
cats = Reference(ws2, min_col=1, min_row=5, max_row=8)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
chart1.width = 14
chart1.height = 8
ws2.add_chart(chart1, "G4")

# Category bar chart
chart2 = BarChart()
chart2.title = "Revenue by category"
chart2.style = 11
chart2.y_axis.title = "Revenue ($)"
data2 = Reference(ws2, min_col=2, min_row=12, max_row=17)
cats2 = Reference(ws2, min_col=1, min_row=13, max_row=17)
chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)
chart2.width = 14
chart2.height = 8
ws2.add_chart(chart2, "G21")

# ---------------- Monthly Trend sheet ----------------
ws3 = wb.create_sheet("Monthly Trend")
ws3["A1"] = "Monthly Revenue Trend"
ws3["A1"].font = TITLE_FONT

months = pd.period_range("2024-10", "2025-09", freq="M")
month_labels = [m.strftime("%b-%y") for m in months]

headers3 = ["Month", "Revenue", "Profit"]
for c, h in enumerate(headers3, start=1):
    cell = ws3.cell(row=3, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL

for i, label in enumerate(month_labels):
    r = 4 + i
    ws3.cell(row=r, column=1, value=label).font = BODY_FONT
    ws3.cell(row=r, column=2, value=f"=SUMIF('Raw Data'!K:K,A{r},'Raw Data'!H:H)").number_format = "$#,##0"
    ws3.cell(row=r, column=3, value=f"=SUMIF('Raw Data'!K:K,A{r},'Raw Data'!J:J)").number_format = "$#,##0"

ws3.column_dimensions["A"].width = 12
ws3.column_dimensions["B"].width = 14
ws3.column_dimensions["C"].width = 14

chart3 = LineChart()
chart3.title = "Monthly revenue trend (Oct 2024 - Sep 2025)"
chart3.y_axis.title = "Revenue ($)"
chart3.x_axis.title = "Month"
data3 = Reference(ws3, min_col=2, min_row=3, max_row=3+len(months))
cats3 = Reference(ws3, min_col=1, min_row=4, max_row=3+len(months))
chart3.add_data(data3, titles_from_data=True)
chart3.set_categories(cats3)
chart3.width = 20
chart3.height = 10
ws3.add_chart(chart3, "E3")

wb.save("/home/claude/projects/retail-sales/retail_sales_performance.xlsx")
print("saved")
